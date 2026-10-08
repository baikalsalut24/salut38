#!/usr/bin/env python3
# Отчёт по одной позиции поставщика: что приходит из pyro_goodlist и pyro_good_params -> xlsx
# Запуск: python3 probe-good.py MC143   (ключи из /etc/bs.env, токен из DATA_DIR)
import os, sys, json, urllib.request, urllib.parse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
art = (sys.argv[1] if len(sys.argv) > 1 else 'MC143').strip()
D = os.environ.get('DATA_DIR', '/var/lib/bs-data')
BASE = os.environ.get('SALUT_BASE', 'https://system.salut-1.ru').rstrip('/')
ID = os.environ.get('SALUT_ID', '')
tk = json.load(open(os.path.join(D, 'supplier-token.json')))['access_token']
def api(path, **kw):
    kw.update(client_id=ID, token=tk)
    with urllib.request.urlopen(BASE + path + '?' + urllib.parse.urlencode(kw), timeout=120) as r:
        return json.loads(r.read().decode())
gl = api('/api/1.0/pyro_goodlist/', empty_count=1)
items = gl.get('data') or gl.get('goods') or gl
if isinstance(items, dict): items = list(items.values())
low = art.lower()
hit = [g for g in items if low in str(g.get('art', '')).lower() or low in str(g.get('name', '')).lower()]
if not hit: sys.exit('Не найдено: ' + art)
g = hit[0]
try: pr = api('/api/1.0/pyro_good_params/', good_id=g['id'])
except Exception as e: pr = {'error': str(e)}
wb = Workbook(); H = Font(bold=True, color='FFFFFF'); F = PatternFill('solid', fgColor='1F3A5F')
def sheet(ws, head, rows, w):
    ws.append(head)
    for c in ws[1]: c.font = H; c.fill = F
    for r in rows: ws.append(r)
    for i, x in enumerate(w): ws.column_dimensions['ABCDEFG'[i]].width = x
    for row in ws.iter_rows(min_row=2):
        for c in row: c.alignment = Alignment(wrap_text=True, vertical='top')
ws = wb.active; ws.title = 'pyro_goodlist'
sheet(ws, ['Поле', 'Значение от поставщика'], [[k, json.dumps(v, ensure_ascii=False) if isinstance(v, (list, dict)) else v] for k, v in g.items()], [22, 80])
ws2 = wb.create_sheet('pyro_good_params')
rows = [[p.get('name'), p.get('value'), p.get('unit')] for p in (pr.get('data') or [])] if isinstance(pr, dict) else []
if not rows: rows = [['(нет данных)', json.dumps(pr, ensure_ascii=False)[:500], '']]
sheet(ws2, ['Параметр', 'Значение', 'Ед.'], rows, [30, 30, 10])
wb.save(f'/tmp/{art}-params.xlsx'); print('Готово: /tmp/%s-params.xlsx' % art)
